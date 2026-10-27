import asyncio
import io
import os
import sys

os.environ.setdefault("DATABASE_URL", "postgresql+asyncpg://u:p@localhost/db")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

import pytest  # noqa: E402
from docx import Document as DocxWriter  # noqa: E402

import main  # noqa: E402
from errors import PermanentIngestError  # noqa: E402
from notebook_chunker import build_document_chunks, embedding_text, split_long_text  # noqa: E402
from notebook_parsing import clean_text, extract_blocks  # noqa: E402


def make_minimal_pdf(text: str) -> bytes:
    """Одностраничный PDF с одной строкой текста (шрифт Helvetica), собранный вручную."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


# --- clean_text -------------------------------------------------------------

def test_clean_text_strips_markdown_link_cookies_and_nbsp():
    assert clean_text("см. [a](http://x) тут") == "см. a тут"
    assert clean_text("Текст\nМы используем файлы cookie, чтобы всё работало") == "Текст"
    assert clean_text("a\xa0b") == "a b"


def test_clean_text_none_and_blank_lines():
    assert clean_text(None) == ""
    assert clean_text("a\n\n\n\nb  \t c") == "a\n\nb c"


# --- split_long_text --------------------------------------------------------

def test_split_long_text_short_text_is_single_part():
    assert split_long_text("Коротко.") == ["Коротко."]


def test_split_long_text_splits_by_sentences_within_limit():
    text = " ".join(f"Предложение номер {i} говорит о важном." for i in range(100))
    parts = split_long_text(text, max_chars=200)
    assert len(parts) > 1
    assert all(len(p) <= 200 for p in parts)
    assert " ".join(parts) == text


def test_split_long_text_hard_cuts_oversized_sentence():
    parts = split_long_text("а" * 3500, max_chars=1400)
    assert [len(p) for p in parts] == [1400, 1400, 700]


# --- build_document_chunks --------------------------------------------------

def block(section, text):
    return {"section": section, "kind": "paragraph", "text": text}


def test_chunks_do_not_merge_different_sections():
    blocks = [block("А", "x" * 100), block("Б", "y" * 100)]
    chunks = build_document_chunks(blocks)
    assert [c["section"] for c in chunks] == ["А", "Б"]
    assert "y" not in chunks[0]["text"] and "x" not in chunks[1]["text"]


def test_chunks_carry_last_short_paragraph_as_overlap():
    p1, p2, p3 = "а" * 900, "б" * 300, "в" * 900
    chunks = build_document_chunks([block("С", p1), block("С", p2), block("С", p3)])
    assert len(chunks) == 2
    assert chunks[0]["text"] == f"{p1}\n{p2}"
    assert chunks[1]["text"] == f"{p2}\n{p3}"


def test_chunks_do_not_carry_long_paragraph_as_overlap():
    p1, p2 = "а" * 1000, "б" * 1000
    chunks = build_document_chunks([block("С", p1), block("С", p2)])
    assert [c["text"] for c in chunks] == [p1, p2]


def test_chunks_shorter_than_80_chars_are_dropped_and_indexes_are_zero_based():
    blocks = [block("А", "коротко"), block("Б", "x" * 100), block("В", "y" * 100)]
    chunks = build_document_chunks(blocks)
    assert [c["chunk_index"] for c in chunks] == [0, 1]
    assert [c["section"] for c in chunks] == ["Б", "В"]


def test_chunks_of_empty_blocks_is_empty():
    assert build_document_chunks([]) == []
    assert build_document_chunks([block("А", "   ")]) == []


# --- extract_blocks ---------------------------------------------------------

def docx_bytes(build) -> bytes:
    doc = DocxWriter()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_blocks_docx_heading_and_table():
    def build(doc):
        doc.add_heading("Стипендия", level=1)
        doc.add_paragraph("Общие положения.")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Колонка"
        table.rows[0].cells[1].text = "Размер"
        table.rows[1].cells[0].text = "значение"
        table.rows[1].cells[1].text = "100"

    blocks = extract_blocks(docx_bytes(build), "docx")
    assert blocks[0] == {"section": "Стипендия", "kind": "heading", "text": "Стипендия"}
    assert blocks[1] == {"section": "Стипендия", "kind": "paragraph", "text": "Общие положения."}
    assert blocks[2] == {
        "section": "Стипендия", "kind": "table", "text": "Колонка: значение | Размер: 100",
    }


def test_extract_blocks_txt_headings_and_paragraphs():
    blocks = extract_blocks("# Раздел\nстрока один\nстрока два\n\nвторой".encode(), "txt")
    assert [b["kind"] for b in blocks] == ["heading", "paragraph", "paragraph"]
    assert blocks[1]["text"] == "строка один строка два"
    assert blocks[1]["section"] == "Раздел"


def test_extract_blocks_txt_ignores_invalid_utf8():
    assert extract_blocks(b"\xff\xfe abc", "txt")[0]["text"].endswith("abc")


def test_extract_blocks_pdf_page_section():
    blocks = extract_blocks(make_minimal_pdf("Hello world"), "pdf")
    assert blocks == [{"section": "Страница 1", "kind": "page", "text": "Hello world"}]


@pytest.mark.parametrize("doc_type", ["docx", "pdf"])
def test_corrupt_file_raises_permanent_error(doc_type):
    with pytest.raises(PermanentIngestError):
        extract_blocks(b"this is not a real file", doc_type)


# --- embedding_text ---------------------------------------------------------

def test_embedding_text_exact_format():
    assert embedding_text("ПУД", "Положения", "Раздел 1", "Текст чанка") == (
        "Название документа: ПУД\nКатегория: Положения\nРаздел: Раздел 1\nТекст: Текст чанка"
    )


# --- chunk_rows / process_one_document --------------------------------------

def test_chunk_rows_fields():
    chunks = [{"chunk_index": 0, "section": "С", "text": "один два три"}]
    rows = main.chunk_rows(chunks)
    assert rows == [{
        "chunk_index": 0, "text": "один два три", "section_path": "С", "start_offset": None,
        "end_offset": None, "token_count": 3, "metadata_json": {"kind": "notebook-v1"},
    }]


def test_embed_batch_sends_passage_kind():
    sent = {}

    class Resp:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *_):
            return False

        def raise_for_status(self):
            pass

        async def json(self):
            return {"embeddings": [[0.0] * main.EMBEDDING_DIM], "model": "m"}

    class Sess:
        def post(self, _url, json):
            sent.update(json)
            return Resp()

    asyncio.run(main._embed_batch(Sess(), ["t"]))
    assert sent["kind"] == "passage"
